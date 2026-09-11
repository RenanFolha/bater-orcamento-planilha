# Desenvolvedor Chefe: RenanFolha

"""Teste de export_service.py: gera a planilha a partir de um registro de
histórico simulado (mesmo formato devolvido por GET /historico/{codigo})
e confere que o resultado é um .xlsx válido e não vazio."""

from io import BytesIO

import openpyxl
import pytest

import export_service as export


def _registro_teste():
    return {
        "codigo": "ORC-0001",
        "criado_em": "2026-01-15T10:30:00",
        "cliente": "Cliente Teste Ltda",
        "responsavel": "Fulano de Tal",
        "origem_resumo": "São Paulo",
        "destino_resumo": "Curitiba, PR, Brasil",
        "veiculo": "VUC",
        "distancia_km": 100,
        "valor_mercadoria": 1000.0,
        "frete_total": 543.21,
        "dados": {
            "payload": {"paletes": [{"comprimento": 40, "largura": 30, "altura": 25}]},
            "resultado": {
                "entrada": {"peso_kg": 50, "transporte": "Rodoviário"},
                "calculos_intermediarios": {
                    "volume_total_m3": 0.03,
                    "peso_considerado_kg": 50,
                    "coleta_terceirizada": False,
                    "entrega_terceirizada": False,
                },
                "resultado": {"prazo_estimado_dias_uteis": 5},
            },
        },
    }


def test_gerar_planilha_orcamento_retorna_xlsx_valido():
    conteudo = export.gerar_planilha_orcamento(_registro_teste())
    assert isinstance(conteudo, bytes)
    assert len(conteudo) > 0

    wb = openpyxl.load_workbook(BytesIO(conteudo))
    ws = wb[wb.sheetnames[0]]
    assert ws.cell(row=6, column=2).value == "ORC-0001"  # B6: ID COTAÇÃO
    assert ws.cell(row=6, column=7).value == 1000.0  # G6: valor da mercadoria
    assert ws.cell(row=6, column=12).value == 543.21  # L6: frete final
    assert ws.cell(row=6, column=8).value == 1  # H6: QTDE (1 linha, quantidade padrão 1)


def test_qtde_soma_a_quantidade_de_cada_linha_de_palete():
    registro = _registro_teste()
    registro["dados"]["payload"]["paletes"] = [
        {"comprimento": 40, "largura": 30, "altura": 25, "quantidade": 3},
        {"comprimento": 50, "largura": 40, "altura": 30, "quantidade": 2},
    ]
    conteudo = export.gerar_planilha_orcamento(registro)
    wb = openpyxl.load_workbook(BytesIO(conteudo))
    ws = wb[wb.sheetnames[0]]
    assert ws.cell(row=6, column=8).value == 5  # H6: QTDE = 3 + 2, não 2 linhas


def test_campos_de_texto_livre_nao_viram_formula_na_planilha():
    # Regressão de segurança: "cliente" e outros campos de texto livre não
    # podem ser gravados como fórmula quando começam com = + - @ (CWE-1236,
    # injeção de fórmula) -- ver export_service._texto_seguro.
    registro = _registro_teste()
    registro["cliente"] = "=cmd|'/c calc'!A0"
    registro["origem_resumo"] = "+SUM(A1:A10)"
    registro["dados"]["resultado"]["entrada"]["transporte"] = "@evil()"
    conteudo = export.gerar_planilha_orcamento(registro)

    wb = openpyxl.load_workbook(BytesIO(conteudo))
    ws = wb[wb.sheetnames[0]]

    # D3 nunca começa com o caractere arriscado (tem o texto fixo "ORÇAMENTO
    # DE FRETE - CLIENTE " sempre na frente), mas ainda assim confere que
    # não virou fórmula e que o texto do cliente aparece íntegro.
    assert ws["D3"].data_type == "s"
    assert "CMD|'/C CALC'!A0" in ws["D3"].value

    celula_origem = ws.cell(row=6, column=3)
    assert celula_origem.data_type == "s"
    assert not celula_origem.value.startswith(("=", "+", "-", "@"))

    celula_transporte = ws.cell(row=6, column=14)
    assert celula_transporte.data_type == "s"
    assert not celula_transporte.value.startswith(("=", "+", "-", "@"))


def test_gerar_planilha_sem_modelo_gera_erro(monkeypatch, tmp_path):
    monkeypatch.setattr(export, "MODELO_PATH", tmp_path / "nao-existe.xlsx")
    with pytest.raises(export.ExportacaoError):
        export.gerar_planilha_orcamento(_registro_teste())
